"""Statistical Centre of Iran (amar.org.ir) consumer price index time series.

The SCI is the body that actually publishes Iran's CPI; the Central Bank's own inflation
pages sit behind an F5 JavaScript challenge (``Request Rejected``) and are not machine
readable, and the IMF's ``dataservices.imf.org`` is unreachable from Iranian networks, so
this is the one working source for a monthly index level.

Shape of the feed (verified 2026-09-20):

* ``https://amar.org.ir/prices`` links one ``/Portals/0/Statistics/ts_<key>_<stamp>.xlsx``
  workbook per series -- ``ts_national`` (whole country), ``ts_urban``, ``ts_rural``. The
  ``<stamp>`` changes on every monthly release, so the newest link has to be discovered
  rather than hardcoded.
* Inside, sheet ``جدول 1`` is the index-level table: row 1 is the title (it names the base
  year, e.g. ``برمبنای سال پایه 100=1400``), row 2 carries the Jalali year on the first
  month of each year only, row 3 repeats the twelve month names, and rows 4+ are COICOP
  groups whose first column is the label -- ``شاخص كل`` being the headline index. Columns
  are pre-allocated past the last release, so trailing ones have no month name.

``amar.org.ir`` serves its certificate **without the intermediate**, so every client that
does not already hold that CA fails with ``unable to get local issuer certificate``. The
workbook requests therefore run with ``verify=False``; the payload is public statistics,
and it is the only way to read the feed at all.
"""

import os
import re
import time
from typing import Dict, List, Optional, Tuple

import jdatetime
import requests

from .._http import get
from ..config import get_config
from ..exceptions import DataUnavailable

__all__ = [
    "BASE",
    "PRICES_URL",
    "DATASETS",
    "HEADLINE_ROW",
    "MAX_CACHE_AGE",
    "dataset_url",
    "workbook",
    "series",
]

BASE = "https://amar.org.ir"
PRICES_URL = f"{BASE}/prices"

#: public alias -> SCI workbook prefix
DATASETS = {
    "CPI": "ts_national",
    "CPI_URBAN": "ts_urban",
    "CPI_RURAL": "ts_rural",
}

HEADLINE_ROW = "شاخص کل"

#: seconds a downloaded workbook is trusted before the prices page is scraped again
MAX_CACHE_AGE = 24 * 3600
INDEX_SHEET = "جدول 1"

_MONTHS = [
    "فروردین",
    "اردیبهشت",
    "خرداد",
    "تیر",
    "مرداد",
    "شهریور",
    "مهر",
    "آبان",
    "آذر",
    "دی",
    "بهمن",
    "اسفند",
]
_MONTH_INDEX = {name: number for number, name in enumerate(_MONTHS, start=1)}
_BASE_YEAR_RE = re.compile(r"100\s*=\s*(\d{4})|(\d{4})\s*=\s*100")


def _normalise(text) -> str:
    if text is None:
        return ""
    return (
        str(text)
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", "")
        .replace("\u200f", "")
        .strip()
    )


def dataset_url(session: requests.Session, prefix: str) -> str:
    """Newest ``/Portals/0/Statistics/<prefix>...xlsx`` link on the prices page."""
    html = get(session, PRICES_URL, expect_json=False, verify=False).text
    pattern = re.compile(
        r'"(/Portals/0/Statistics/(%s)[-_][^"]*\.xlsx)"' % re.escape(prefix), re.IGNORECASE
    )
    matches = [match.group(1) for match in pattern.finditer(html)]
    if not matches:
        raise DataUnavailable(404, PRICES_URL, f"no {prefix!r} workbook linked from {PRICES_URL}")
    # the trailing stamp is a Jalali publication timestamp, so lexical max == newest
    return BASE + max(matches)


def workbook(session: requests.Session, prefix: str, *, refresh: bool = False) -> str:
    """Newest workbook for ``prefix``, downloaded at most once per ``MAX_CACHE_AGE``.

    SCI releases one workbook a month, and the prices page that has to be scraped to find
    it is ~600 KB, so a cached copy younger than a day short-circuits both requests.
    """
    directory = os.path.join(get_config().resolved_cache_dir(), "sci")
    os.makedirs(directory, exist_ok=True)
    if not refresh:
        cached = _freshest_cached(directory, prefix)
        if cached is not None:
            return cached
    url = dataset_url(session, prefix)
    path = os.path.join(directory, os.path.basename(url))
    if refresh or not os.path.exists(path):
        response = get(session, url, expect_json=False, verify=False, timeout=120)
        with open(path, "wb") as handle:
            handle.write(response.content)
    return path


def _freshest_cached(directory: str, prefix: str) -> Optional[str]:
    candidates = [
        os.path.join(directory, name)
        for name in os.listdir(directory)
        if name.lower().startswith(prefix.lower()) and name.lower().endswith(".xlsx")
    ]
    if not candidates:
        return None
    newest = max(candidates, key=os.path.getmtime)
    return newest if (time.time() - os.path.getmtime(newest)) < MAX_CACHE_AGE else None


def _load_rows(path: str) -> List[tuple]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dependency
        raise DataUnavailable(0, path, "reading SCI workbooks requires openpyxl") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if INDEX_SHEET not in book.sheetnames:
            raise DataUnavailable(0, path, f"{path} has no {INDEX_SHEET!r} sheet")
        return list(book[INDEX_SHEET].iter_rows(values_only=True))
    finally:
        book.close()


_parsed: Dict[Tuple[str, str], Tuple[List[Tuple], Dict]] = {}


def parse_workbook(path: str, row_label: str = HEADLINE_ROW) -> Tuple[List[Tuple], Dict]:
    """``([(date, value), ...], meta)`` for one COICOP row, oldest first."""
    memo_key = (path, row_label)
    if memo_key in _parsed:
        return _parsed[memo_key]
    rows = _load_rows(path)
    if len(rows) < 4:
        raise DataUnavailable(0, path, f"{path} has no data rows")

    title = _normalise(rows[0][0])
    years, months = rows[1], rows[2]
    wanted = _normalise(row_label)
    target: Optional[tuple] = None
    for row in rows[3:]:
        if _normalise(row[0]) == wanted:
            target = row
            break
    if target is None:
        raise DataUnavailable(0, path, f"{path} has no {row_label!r} row")

    observations: List[Tuple] = []
    year: Optional[int] = None
    for column in range(1, len(months)):
        if column < len(years) and years[column] not in (None, ""):
            try:
                year = int(str(years[column]).strip())
            except ValueError:
                year = None
        month = _MONTH_INDEX.get(_normalise(months[column]))
        if year is None or month is None or column >= len(target):
            continue
        value = target[column]
        try:
            number = float(value)
        except (TypeError, ValueError):
            continue
        observations.append((jdatetime.date(year, month, 1).togregorian(), number))

    observations.sort(key=lambda item: item[0])
    match = _BASE_YEAR_RE.search(title)
    base_year = next((int(group) for group in (match.groups() if match else ()) if group), None)
    result = (observations, {"title": title, "base_year": base_year, "row": wanted, "path": path})
    _parsed[memo_key] = result
    return result


def series(
    session: requests.Session,
    alias: str,
    *,
    row_label: str = HEADLINE_ROW,
    refresh: bool = False,
) -> Tuple[List[Tuple], Dict]:
    prefix = DATASETS.get(alias.upper())
    if prefix is None:
        raise DataUnavailable(0, PRICES_URL, f"unknown SCI series: {alias!r}")
    return parse_workbook(workbook(session, prefix, refresh=refresh), row_label)
