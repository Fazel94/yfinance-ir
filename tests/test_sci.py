import os
import time

import openpyxl
import pandas as pd
import pytest

from yfinance_ir.config import get_config
from yfinance_ir.exceptions import DataUnavailable
from yfinance_ir.resolver import resolve
from yfinance_ir.sources import sci
from yfinance_ir.ticker import Ticker

PRICES_HTML = """
<html><body>
<a href="/Portals/0/Statistics/ts_national_140401-14040101090000.xlsx">old</a>
<a href="/Portals/0/Statistics/ts_national_140505-14050618165053.xlsx">new</a>
<a href="/Portals/0/Statistics/ts_urban_140505-14050618165804.xlsx">urban</a>
<a href="/Portals/0/Statistics/ts_ppi_140404-14050407160106.xlsx">ppi</a>
</body></html>
"""

MONTHS = ["فروردین", "اردیبهشت ", "خرداد", "تیر", "مرداد", "شهریور",
          "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]


def _workbook(tmp_path, levels, *, title="شاخص قیمت مصرف کننده ... برمبنای سال پایه 100=1400",
              pad=2):
    """Build a workbook shaped exactly like SCI's ``جدول 1``.

    ``levels`` maps a Jalali year to its twelve monthly values (``None`` = not published).
    """
    book = openpyxl.Workbook()
    book.remove(book.active)
    book.create_sheet("فهرست")
    sheet = book.create_sheet("جدول 1")
    sheet.cell(row=1, column=1, value=title)
    column = 2
    for year, values in levels.items():
        for offset, (month, value) in enumerate(zip(MONTHS, values)):
            if offset == 0:  # SCI writes the year only above the first month
                sheet.cell(row=2, column=column, value=year)
            sheet.cell(row=3, column=column, value=month)
            if value is not None:
                sheet.cell(row=4, column=column, value=value)
                sheet.cell(row=5, column=column, value=value * 2)
            column += 1
    for _ in range(pad):  # SCI pre-allocates columns with no month name
        sheet.cell(row=4, column=column, value=None)
        column += 1
    sheet.cell(row=4, column=1, value="شاخص كل")  # Arabic kaf, as filed
    sheet.cell(row=5, column=1, value="خوراكی\u200cها")
    path = str(tmp_path / "ts_national_140505-14050618165053.xlsx")
    book.save(path)
    return path


def test_parse_reads_month_columns_year_by_year(tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12, 1404: [110.0] * 6 + [None] * 6})

    observations, meta = sci.parse_workbook(path)

    assert len(observations) == 18  # 12 + 6 published months, padding ignored
    assert observations[0][0] == pd.Timestamp("2024-03-20").date()  # 1403-01-01
    assert observations[12][1] == 110.0
    assert meta["base_year"] == 1400
    assert meta["row"] == "شاخص کل"


def test_the_headline_row_is_matched_across_arabic_spelling(tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12})

    observations, _meta = sci.parse_workbook(path, row_label="شاخص کل")  # Persian kaf

    assert observations[0][1] == 100.0


def test_another_coicop_row_can_be_selected(tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12})

    observations, meta = sci.parse_workbook(path, row_label="خوراکی\u200cها")

    assert observations[0][1] == 200.0
    assert meta["row"] == "خوراکیها"


def test_a_missing_row_is_reported(tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12})

    with pytest.raises(DataUnavailable, match="ندارد"):
        sci.parse_workbook(path, row_label="ندارد")


def test_dataset_url_picks_the_newest_publication_stamp(api, session):
    api.mock.add("GET", sci.PRICES_URL, body=PRICES_HTML, content_type="text/html")

    url = sci.dataset_url(session, "ts_national")

    assert url.endswith("ts_national_140505-14050618165053.xlsx")


def test_an_unlinked_dataset_raises(api, session):
    api.mock.add("GET", sci.PRICES_URL, body="<html></html>", content_type="text/html")

    with pytest.raises(DataUnavailable, match="ts_rural"):
        sci.dataset_url(session, "ts_rural")


def test_a_fresh_cached_workbook_skips_both_requests(api, session, tmp_path):
    directory = os.path.join(get_config().resolved_cache_dir(), "sci")
    os.makedirs(directory, exist_ok=True)
    cached = os.path.join(directory, "ts_national_140505-14050618165053.xlsx")
    open(cached, "wb").close()

    assert sci.workbook(session, "ts_national") == cached
    assert not api.mock.calls  # no scrape, no download


def test_a_stale_cached_workbook_is_replaced(api, session):
    directory = os.path.join(get_config().resolved_cache_dir(), "sci")
    os.makedirs(directory, exist_ok=True)
    stale = os.path.join(directory, "ts_national_140401-14040101090000.xlsx")
    open(stale, "wb").close()
    os.utime(stale, (0, time.time() - sci.MAX_CACHE_AGE - 60))
    api.mock.add("GET", sci.PRICES_URL, body=PRICES_HTML, content_type="text/html")
    api.mock.add(
        "GET",
        sci.BASE + "/Portals/0/Statistics/ts_national_140505-14050618165053.xlsx",
        body=b"xlsx-bytes",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    path = sci.workbook(session, "ts_national")

    assert path.endswith("ts_national_140505-14050618165053.xlsx")
    assert open(path, "rb").read() == b"xlsx-bytes"


def test_cpi_resolves_without_touching_the_network(session):
    instrument = resolve(session, "cpi")

    assert (instrument.source, instrument.kind) == ("sci", "MACRO")
    assert instrument.ins_code == "CPI"


def test_history_is_monthly_with_computed_inflation(monkeypatch, tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12, 1404: [110.0, 121.0] + [None] * 10})
    monkeypatch.setattr(sci, "workbook", lambda *a, **k: path)

    frame = Ticker("CPI").history(period="max")

    assert len(frame) == 14
    assert list(frame.index.month) == [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3, 4]
    row = frame.iloc[13]
    assert row["Open"] == row["High"] == row["Low"] == row["Close"] == 121.0
    assert pd.isna(row["Volume"])
    assert row["MoM"] == pytest.approx(10.0)  # 110 -> 121
    assert frame["YoY"].iloc[12] == pytest.approx(10.0)  # 100 -> 110, twelve months apart
    assert pd.isna(frame["MoM"].iloc[0])


def test_info_reports_the_base_year_and_latest_inflation(monkeypatch, tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12, 1404: [110.0] + [None] * 11})
    monkeypatch.setattr(sci, "workbook", lambda *a, **k: path)

    info = Ticker("CPI").info

    assert info["quoteType"] == "MACRO"
    assert info["frequency"] == "monthly"
    assert info["baseYear"] == 1400
    assert info["dataset"] == "ts_national"
    assert info["annualInflation"] == pytest.approx(10.0)
    assert info["currency"] is None


def test_only_monthly_intervals_are_accepted(monkeypatch, tmp_path):
    path = _workbook(tmp_path, {1403: [100.0] * 12})
    monkeypatch.setattr(sci, "workbook", lambda *a, **k: path)
    ticker = Ticker("CPI")

    assert len(ticker.history(interval="1mo")) == 12
    with pytest.raises(NotImplementedError):
        ticker.history(interval="1h")


def test_an_unknown_series_alias_is_rejected(session):
    with pytest.raises(DataUnavailable, match="unknown SCI series"):
        sci.series(session, "CPI_MARS")
